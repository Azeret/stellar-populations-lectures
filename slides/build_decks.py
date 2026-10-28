"""Builds the two lecture decks as .pptx in the slide style of the IMF talks
(dark navy background, blue bold titles, white text), with animations,
sketches, Slido slides and speaker notes.  Run: python build_decks.py
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
ANIM = os.path.join(ROOT, "animations", "out")
SKETCH = os.path.join(ROOT, "figures", "sketches")
OUT = HERE

NAVY = RGBColor(0x1B, 0x1A, 0x2E); BLUE = RGBColor(0x1E, 0x8F, 0xD5); WHITE = RGBColor(0xF4, 0xF4, 0xF6)
GREY = RGBColor(0x9A, 0x9A, 0xB5); ORANGE = RGBColor(0xF5, 0xA6, 0x23); PINK = RGBColor(0xEF, 0x5D, 0x8F)
GREEN = RGBColor(0x5C, 0xC8, 0xA1)
FONT = "Montserrat"
W, H = Inches(13.333), Inches(7.5)
SLIDO_CODES = {"L1": "#2458697", "L2": "#1642685"}   # Slido event codes


class Deck:
    def __init__(self, title):
        self.p = Presentation(); self.p.slide_width, self.p.slide_height = W, H
        self.title = title; self.n = 0
        self.slido_footer = f"slido.com  \u00b7  {SLIDO_CODES.get(title, '#CODE')}"

    # -------------------------------------------------------------- helpers
    def _slide(self, notes=""):
        s = self.p.slides.add_slide(self.p.slide_layouts[6])
        bg = s.background.fill; bg.solid(); bg.fore_color.rgb = NAVY
        if notes:
            s.notes_slide.notes_text_frame.text = notes.strip()
        self.n += 1
        return s

    def _text(self, s, x, y, w, h, text, size=24, color=WHITE, bold=False, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP, italic=False):
        tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
        tf.vertical_anchor = anchor
        lines = text if isinstance(text, list) else [text]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run(); r.text = line
            r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold; r.font.italic = italic
            r.font.name = FONT
            p.space_after = Pt(size * 0.45)
        return tb

    def _title(self, s, text, sub=None):
        self._text(s, Inches(0.45), Inches(0.25), Inches(12.4), Inches(0.9), text, size=34, color=BLUE, bold=True)
        if sub:
            self._text(s, Inches(0.45), Inches(1.0), Inches(12.4), Inches(0.6), sub, size=20, color=WHITE, bold=True)

    def _footer(self, s):
        self._text(s, Inches(12.3), Inches(7.0), Inches(0.9), Inches(0.4), str(self.n), size=12, color=GREY,
                   align=PP_ALIGN.RIGHT)

    def _image(self, s, path, x, y, w=None, h=None):
        if os.path.exists(path):
            return s.shapes.add_picture(path, x, y, w, h)
        return self._text(s, x, y, w or Inches(5), h or Inches(1), f"[missing: {os.path.basename(path)}]", 16, GREY)

    def _video(self, s, name, x, y, w, h):
        mp4 = os.path.join(ANIM, name + ".mp4"); poster = os.path.join(ANIM, name + "_final.png")
        if os.path.exists(mp4) and os.path.getsize(mp4) > 10_000:
            return s.shapes.add_movie(mp4, x, y, w, h, poster_frame_image=poster if os.path.exists(poster) else None,
                                      mime_type="video/mp4")
        return self._image(s, poster, x, y, w, h)

    # -------------------------------------------------------------- slide kinds
    def title_slide(self, title, subtitle, notes=""):
        s = self._slide(notes)
        self._text(s, Inches(0.8), Inches(1.6), Inches(11.7), Inches(2.2), title, 48, BLUE, True, anchor=MSO_ANCHOR.BOTTOM)
        self._text(s, Inches(0.8), Inches(4.0), Inches(11.7), Inches(1.0), subtitle, 24, WHITE)
        self._text(s, Inches(0.8), Inches(5.6), Inches(11.7), Inches(1.0),
                   ["Dr. Tereza Jeřábková", "Galaxies course · Masaryk University · ÚTFA"], 18, GREY)
        return s

    def section(self, number, title, question, notes=""):
        s = self._slide(notes)
        self._text(s, Inches(0.8), Inches(2.3), Inches(2), Inches(1.5), str(number), 96, BLUE, True)
        self._text(s, Inches(2.4), Inches(2.5), Inches(10.4), Inches(1.0), title, 40, WHITE, True)
        self._text(s, Inches(2.4), Inches(3.6), Inches(10.4), Inches(1.0), question, 24, GREY, italic=True)
        return s

    def bullets(self, title, items, image=None, sub=None, notes="", img_w=6.2):
        s = self._slide(notes); self._title(s, title, sub)
        top = Inches(1.7 if sub else 1.4)
        tw = Inches(12.3 - (img_w + 0.3 if image else 0))
        self._text(s, Inches(0.6), top, tw, Inches(5.4), items, 24, WHITE)
        if image:
            self._image(s, image, Inches(13.0 - img_w), top, Inches(img_w))
        self._footer(s); return s

    def figure(self, title, image, caption=None, notes="", sub=None):
        s = self._slide(notes); self._title(s, title, sub)
        self._image(s, image, Inches(1.4), Inches(1.35), Inches(10.5))
        if caption:
            self._text(s, Inches(0.6), Inches(6.85), Inches(12), Inches(0.5), caption, 16, GREY)
        self._footer(s); return s

    def video(self, title, name, caption=None, notes=""):
        s = self._slide(notes)      # the animation carries its own title; `title` is kept for the outline only
        self._video(s, name, Inches(0.4), Inches(0.2), Inches(12.53), Inches(7.05))
        if caption:
            s.notes_slide.notes_text_frame.text += "\n\nCaption: " + caption
        return s

    def slido(self, code, kind, question, options=None, notes=""):
        s = self._slide(notes)
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(0.5), Inches(4.2), Inches(0.6))
        box.fill.solid(); box.fill.fore_color.rgb = GREEN; box.line.fill.background()
        tf = box.text_frame; tf.text = f"Slido {code} · {kind}"
        for r in tf.paragraphs[0].runs:
            r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = NAVY; r.font.name = FONT
        self._text(s, Inches(0.6), Inches(1.6), Inches(12.1), Inches(2.4), question, 40, WHITE, True)
        if options:
            self._text(s, Inches(0.9), Inches(3.9), Inches(11.6), Inches(3.0),
                       [f"{chr(65 + i)}   {o}" for i, o in enumerate(options)], 26, WHITE)
        self._text(s, Inches(0.6), Inches(6.8), Inches(12.1), Inches(0.5), self.slido_footer, 18, GREEN, True, PP_ALIGN.RIGHT)
        return s

    def insert_marker(self, what, source, notes=""):
        s = self._slide(notes)
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(2.2), Inches(10.3), Inches(3.1))
        box.fill.background(); box.line.color.rgb = ORANGE; box.line.width = Pt(3); box.line.dash_style = 4
        self._text(s, Inches(1.9), Inches(2.5), Inches(9.5), Inches(2.6),
                   ["→ insert your slides here", what, f"from: {source}"], 28, ORANGE, True, PP_ALIGN.CENTER)
        return s

    def big_statement(self, text, sub=None, notes="", color=WHITE):
        s = self._slide(notes)
        self._text(s, Inches(1.0), Inches(2.3), Inches(11.3), Inches(2.0), text, 40, color, True, PP_ALIGN.CENTER)
        if sub:
            self._text(s, Inches(1.0), Inches(4.4), Inches(11.3), Inches(1.5), sub, 24, GREY, False, PP_ALIGN.CENTER)
        return s

    def save(self, name):
        path = os.path.join(OUT, name); self.p.save(path); print("saved", path, self.n, "slides")


# =============================================================================
# LECTURE 1
# =============================================================================
def lecture1():
    d = Deck("L1")
    d.title_slide("Stellar populations & the stellar IMF",
                  "Lecture 1 — How do you weigh a galaxy you cannot resolve?",
                  notes="0:00. Welcome. Put the Slido code on the board now; everybody joins before slide 3.")
    d.slido("S1.1", "open text", "Andromeda: how much does it weigh, and how would you find out?",
            notes="0:02. Show your Andromeda image (ISTA p.32) behind this if you like. Collect ideas: rotation "
                  "curve, light × M/L, lensing, satellites. Keep the list: we come back to it at the end.")
    d.slido("S1.2", "word cloud", "What can we actually measure from a galaxy we cannot resolve?",
            notes="0:06. Expected: colours, magnitudes, spectra, line widths, sizes. Punchline: only light. "
                  "2 minutes pair talk first, then the word cloud.")
    d.figure("Today's story: from light to mass", os.path.join(SKETCH, "s6_light_to_mass.png"),
             caption="Everything we say about a galaxy's mass, age or SFR goes through a model — and every model has an IMF.",
             notes="0:12. The roadmap. Lecture 1 builds the M/L box; Lecture 2 opens the chemistry box. "
                   "The IMF hides inside both.")
    d.section(1, "The IMF: counting stars like dogs", "What does it mean that “the IMF has a slope of 2.35”?",
              notes="0:15")
    d.insert_marker("the dog analogy: ‘Let's study dogs like we study stars’ → dog IMFs in different environments",
                    "Canva ‘ISTA seminar’, slides 21–30",
                    notes="0:15–0:28. Bike ride with a camera: distance, breed (= type), age, multiplicity, cats "
                          "(contaminants). Then the dog IMF and the city/village dog IMFs. Keep the house-dog joke.")
    d.slido("S1.3", "multiple choice", "Counting dogs from a bike: which bias distorts the dog IMF most?",
            ["far-away dogs are missed", "breeds look alike", "puppies vs adults", "dogs walk in pairs", "cats"],
            notes="0:28. No single right answer; map each to a stellar correction: Malmquist/completeness, "
                  "spectral type, age/evolution, unresolved binaries, contaminants.")
    d.insert_marker("Salpeter 1955 and the corrections from luminosity function to IMF; ‘the IMF is a Hilfskonstrukt’",
                    "Canva ‘ISTA seminar’, slides 20 and 31–39",
                    notes="0:30–0:40. The PDLF → IMF chain: mass–luminosity relation, scale height vs age, "
                          "MS brightening, evolved stars, metallicity, unresolved binaries, extinction. Scalo (1998) quote.")
    d.figure("Three IMFs you will meet in papers", os.path.join(SKETCH, "s7_imf_shapes.png"),
             caption="ξ(m) = dN/dm ∝ m^−α.  Salpeter: α = 2.35.  Kroupa: 1.3 below 0.5 M☉, 2.3 above.  Chabrier: log-normal + power law.",
             notes="0:40. They agree above ~1 Msun and differ where stars are faint. Mention that 'canonical' "
                   "means 'what we measured in the Solar neighbourhood', not 'proven universal'.")
    d.slido("S1.4", "quiz", "Salpeter's slope α in dN/dm ∝ m^−α is …",
            ["1.35", "2.35", "3.5", "0.5"],
            notes="0:43. Answer 2.35. Trap: 1.35 is the same IMF written per log m (dN/dlog m ∝ m^−1.35).")
    d.bullets("Game 1 — build a star cluster",
              ["In pairs, draw 30 stars: a random number 000–999 → look up the mass.",
               "000–431  0.13 M☉   ·   432–759  0.32   ·   760–901  0.68",
               "902–959  1.4   ·   960–992  3.5   ·   993–997  12   ·   998–999  42",
               "Add up your cluster mass. Note your most massive star.",
               "Send both to Slido. Did anybody get a supernova progenitor (> 8 M☉)?"],
              notes="0:45–0:58. Random numbers: phone (e.g. 'random number 0 999'), or three d10. "
                    "Expect ~18 Msun per pair and roughly one pair in five with a star > 8 Msun.")
    d.slido("S1.5", "open text", "Your cluster: total mass and most massive star  (e.g. “19 M☉, 3.5 M☉”)",
            notes="0:52. Read some out. Then add up the room as one cluster. Discussion: bigger clusters host "
                  "bigger stars — random sampling or a physical m_max–M_cl relation? (open debate).")
    d.video("Drawing stars from the IMF", "a1_imf_sampling",
            "From 10 to 100 000 stars: the histogram converges to the IMF; massive stars stay rare.",
            notes="0:55. Pause at ~300 stars (first massive stars appear) and at the end (617 stars > 8 Msun out of 100 000).")
    d.slido("S1.6", "quiz", "What fraction of newborn stars are less massive than the Sun?  …and what fraction of the mass goes into stars above 8 M☉?",
            ["50%  /  2%", "75%  /  50%", "90%  /  ~20%", "99%  /  80%"],
            notes="0:58. Answer C: 90% of stars are < 1 Msun; 23% of the mass formed is in stars > 8 Msun (Kroupa).")
    d.figure("Number, mass and light live in different stars", os.path.join(SKETCH, "s5_number_mass_light.png"),
             caption="Kroupa IMF at birth. Light: main-sequence L(m). After ~1 Gyr the light comes from red giants instead.",
             notes="1:00. The key slide of part 1. Low-mass stars: number. Intermediate: mass. Massive: light (when young). "
                   "So the light tells you about a few stars, and you extrapolate to all the mass with the IMF.")
    d.big_statement("10 minutes break", "Puzzle for the break: a galaxy's light is 99% from <1% of its stars. How would you weigh it?",
                    notes="1:05–1:15")
    d.section(2, "Building a population", "How does a population of stars look from far away?", notes="1:15")
    d.figure("The recipe of a simple stellar population (SSP)", os.path.join(SKETCH, "s1_ssp_recipe.png"),
             caption="Fλ(t, Z) = ∫ ξ(m) fλ(m, t, Z) dm  over the stars still alive",
             notes="1:15. One age, one metallicity. Real galaxies = sums of SSPs (star-formation history).")
    d.bullets("FaSTAR: SSPs at any age, metallicity and IMF",
              ["Martín-Navarro et al. 2026, A&A 711, A282 & A283",
               "Neural-network stellar spectra + isochrones + any IMF you like",
               "Fast and differentiable: ages 20 Myr–14 Gyr, [M/H] −2.5 to +0.3",
               "Integrated (infinitely many stars) and semi-resolved (N stars) modes",
               "pip install fastar-astro  →  notebook 02"],
             notes="1:20. Everything in the next animations is computed with FaSTAR in notebook 02; students can redo it.")
    d.slido("S1.7", "multiple choice", "Same mass formed: how much fainter (V band) is a 10 Gyr population than a 100 Myr one?",
            ["3×", "10×", "~35×", "1000×"],
            notes="1:23. Answer ~35× (FaSTAR, Kroupa, solar). Then play A2.")
    d.video("One population, getting older", "a2_ssp_ageing",
            "FaSTAR, Kroupa IMF, solar metallicity. M/L_V (stars + remnants) grows ~70× from 30 Myr to 13 Gyr.",
            notes="1:25. Watch the Balmer lines grow (A stars, ~0.3–1 Gyr) and fade. Blue light disappears as massive stars die.")
    d.slido("S1.8", "multiple choice", "A galaxy's spectrum is redder than your model. The galaxy is …",
            ["older", "more metal-rich", "dustier", "can't tell from colour alone"],
            notes="1:32. Answer D — the age–metallicity(–dust) degeneracy. Then A3.")
    d.video("Same age, more metals", "a3_ssp_metallicity",
            "FaSTAR, 10 Gyr. Metal lines deepen and the continuum reddens — like getting older (Worthey's 3/2 rule).",
            notes="1:34. Point at Mg b and Fe lines: index ratios break the degeneracy (Lick indices).")
    d.slido("S1.9", "multiple choice", "At 10 Gyr, which IMF has a higher M/L than Salpeter?",
            ["top-heavy only (α = 1.3)", "bottom-heavy only (α = 3)", "both", "neither"],
            notes="1:42. Answer C. Then A4.")
    d.video("Change the IMF: the spectrum barely moves, the mass does", "a4_ssp_imf",
            "Single power law 0.1–100 M☉, 10 Gyr, solar. Top-heavy: mass in remnants. Bottom-heavy: mass in faint dwarfs.",
            notes="1:44. The spectrum changes by a few per cent (Na D, TiO: dwarf-sensitive), M/L by a factor ~4. "
                  "This is why IMF measurements in galaxies are hard and controversial.")
    d.slido("S1.10", "multiple choice", "Two clusters with 300 stars each, same age, metallicity and IMF. Same spectrum?",
            ["yes", "no"], notes="1:52. Answer no. Then A5.")
    d.video("Few stars, noisy light", "a5_stochastic_0.1Gyr",
            "FaSTAR semi-resolved models, 100 Myr: a handful of massive or evolved stars decides the light.",
            notes="1:54. Relevant for star clusters, dwarf galaxies and for ELT-resolution data (semi-resolved regime).")
    d.section(3, "So what?", "Every stellar mass and SFR you read carries an IMF assumption", notes="2:02")
    d.insert_marker("Hα SFR and stellar masses with a variable IMF (Haslbauer, Yan, Jeřábková 2024)",
                    "Canva ‘ISTA seminar’, slides 47–48",
                    notes="2:03. Hα counts ionising photons = massive stars only. To get an SFR you extrapolate with the IMF.")
    d.slido("S1.11", "quiz", "If the true IMF is top-heavy, an Hα SFR computed assuming Kroupa is …",
            ["too high", "too low", "unchanged"],
            notes="2:10. Answer: too high — more ionising photons per Msun formed than Kroupa assumes.")
    d.bullets("Take-aways",
              ["The IMF is inferred, not observed — a useful Hilfskonstrukt.",
               "Most stars are small; the light comes from a few massive or evolved ones.",
               "M/L depends on age, metallicity — and on the IMF (U-shape).",
               "So every galaxy mass or SFR carries an IMF assumption.",
               "Next time: chemistry as a clock — and the IMF again."],
              notes="2:14")
    d.bullets("Homework & exit ticket",
              ["Notebooks 01 (IMF) and 02 (SSPs with FaSTAR)",
               "github.com/Azeret/stellar-populations-lectures   [QR code]",
               "Slido S1.12: one thing still unclear? How was the pace?"],
              notes="2:17. Put the QR code here in Canva.")
    d.save("lecture1_stellar_populations_IMF.pptx")


# =============================================================================
# LECTURE 2
# =============================================================================
def lecture2():
    d = Deck("L2")
    d.title_slide("Chemical evolution & the [Mg/Fe] clock",
                  "Lecture 2 — How do you read a galaxy's clock from its chemistry?",
                  notes="0:00. Slido code on the board.")
    d.slido("S2.1", "quiz", "Recap: in which year did Salpeter publish the IMF?  Where is most of the stellar mass?  How does M/L depend on the IMF slope?",
            ["1955 · low-mass stars · U-shape", "1975 · massive stars · rising", "1955 · massive stars · falling"],
            notes="0:01. Three-question quiz in Slido (set up as three separate questions).")
    d.slido("S2.2", "word cloud", "Name one element in your body — and where was it made?",
            notes="0:07. Expect: Big Bang (H), stars, supernovae. Then reveal the periodic table of origins.")
    d.insert_marker("the periodic table of element origins",
                    "Johnson (2019, Science 363, 474) / NASA, CC BY-SA — download and place here",
                    notes="0:10. Big Bang (H, He, Li), dying low-mass stars (AGB: C, N, s-process), exploding massive stars "
                          "(O, Mg, Si, Ca), exploding white dwarfs (Fe, Ni), merging neutron stars (r-process).")
    d.bullets("Notation", ["[X/H] = log₁₀(N_X/N_H) − log₁₀(N_X/N_H)☉",
                           "[Mg/Fe] = [Mg/H] − [Fe/H]",
                           "[Fe/H] = −1 → one tenth of the solar iron per hydrogen",
                           "Z = mass fraction of all metals (Sun: Z ≈ 0.013–0.014, Asplund+2009)",
                           "[α/Fe] > 0 → more O, Mg, Si, Ca per Fe than the Sun"],
              notes="0:17")
    d.slido("S2.3", "quiz", "A star has 1/10 of the solar Fe/H. Its [Fe/H] is …", ["−1", "−0.1", "0.1", "−10"],
            notes="0:24. Answer −1.")
    d.section(1, "The galaxy as a bathtub", "What goes in, what comes out, what stays?", notes="0:27")
    d.figure("The one-zone (bathtub) model", os.path.join(SKETCH, "s2_bathtub.png"),
             caption="dM_gas/dt = inflow − (1−R)·SFR − η·SFR      dM_X/dt = −(1−R+η)·SFR·X_gas + yields(SN II, SN Ia, AGB)",
             notes="0:27–0:40. Draw it on the board together. Closed box: Z = y·ln(1/μ) → too many metal-poor "
                   "G dwarfs (G-dwarf problem) → we need inflow. One sentence each.")
    d.section(2, "Two kinds of supernovae", "Why is the ratio of magnesium to iron a clock?", notes="0:40")
    d.figure("Prompt and delayed enrichment", os.path.join(SKETCH, "s3_sn_timeline.png"),
             caption="SN II: within ~40 Myr, Mg-rich.  SN Ia: delay-time distribution ~1/t from ~40 Myr to a Hubble time, Fe-rich.",
             notes="0:40. Before the game, so students know the rules of nature they are playing.")
    d.bullets("Game 2 — the cosmic kitchen",
              ["Your team is a galaxy with 8 units of star formation for 8 rounds (1 round ≈ 250 Myr).",
               "Each round: (1) add the Fe from SN Ia due now; (2) form stars: +3 Mg and +1 Fe per unit (SN II);",
               "book SN Ia: +1 Fe next round, +1 Fe the round after, +1 Fe four rounds later;",
               "(3) write down Mg/Fe in your pot → [Mg/Fe] = log(Mg/Fe).",
               "Fast: 4 + 4 units in rounds 1–2.  Steady: 1 unit per round.  Free: beat Fast!"],
              notes="0:43–0:55. Hand out the tally sheet. Log table: 3→+0.48, 2→+0.30, 1.5→+0.18, 1.2→+0.08, 1→0, 0.75→−0.12.")
    d.slido("S2.4", "open text", "Your team's mass-weighted [Mg/Fe] of the stars you formed",
            notes="0:53. Fast ≈ +0.40, Steady ≈ +0.18; all 8 units in round 1 gives +0.48 (the winner).")
    d.figure("The plateau and the knee", os.path.join(SKETCH, "s4_knee_schematic.png"),
             caption="Plateau: SN II only.  Knee: SN Ia catch up.  Its [Fe/H] tells how fast metals built up before ~1 Gyr.",
             notes="0:56")
    d.slido("S2.5", "multiple choice", "A dwarf galaxy forms stars slowly. Its [Mg/Fe] knee sits at …",
            ["lower [Fe/H]", "the same [Fe/H]", "higher [Fe/H]"], notes="1:00. Answer: lower [Fe/H].")
    d.big_statement("10 minutes break", "Puzzle: could a galaxy have [Mg/Fe] above the SN II plateau?", notes="1:05–1:15")
    d.video("[Mg/Fe] is a clock", "a6_mgfe_clock",
            "Toy one-zone model (notebook 03): elliptical (fast, quenched at 0.5 Gyr), Milky-Way-like disc, slow leaky dwarf.",
            notes="1:15. The elliptical's stars stay α-rich because star formation stopped before SN Ia caught up.")
    d.slido("S2.6", "multiple choice", "Make the IMF top-heavy. The [Mg/Fe] plateau …", ["goes up", "goes down", "stays"],
            notes="1:22. Answer: up — more of the most massive SN II, which make the most Mg per Fe.")
    d.video("A different IMF moves the plateau", "a7_mgfe_imf",
            "Same Milky-Way-like model, high-mass IMF slope from 1.9 (top-heavy) to 2.7 (top-light).",
            notes="1:24. The degeneracy: high [Mg/Fe] = fast star formation OR top-heavy IMF.")
    d.section(3, "Real models: galevo", "Does it hold up with real stellar lifetimes and yields?", notes="1:28")
    d.insert_marker("GalIMF / galevo: modular open-source code, IGIMF theory",
                    "Canva ‘ISTA seminar’, slides 42–45 (and 48: galaxy-wide IMF variations)",
                    notes="1:28. What galevo adds: lifetimes, metallicity-dependent yields, many elements, SN Ia DTD, "
                          "and a galaxy-wide IMF from the IGIMF theory. github.com/Azeret/galIMF")
    d.video("Same mass, different tempo (galevo)", "a8_galevo_tracks",
            "galevo, 3×10⁹ M☉ formed with a flat SFH over 30 Myr, 300 Myr, 3 Gyr; Kroupa IMF vs IGIMF.",
            notes="1:33")
    d.slido("S2.7", "ranking", "Rank by the [Mg/Fe] of their stars: formed in 30 Myr / 300 Myr / 3 Gyr",
            notes="1:38. Answer: 30 Myr > 300 Myr > 3 Gyr.")
    d.insert_marker("observations: Milky Way thick vs thin disc (APOGEE), dwarf spheroidals with a low knee (Tolstoy+2009)",
                    "literature figures — add your favourites",
                    notes="1:40. Use S2.8 here: 'Track A with a knee at [Fe/H] ≈ −2 belongs to…' (dwarf spheroidal).")
    d.slido("S2.8", "multiple choice", "Track A has its knee at [Fe/H] ≈ −2. It belongs to …",
            ["a dwarf spheroidal", "the Milky Way thick disc", "a giant elliptical"], notes="1:45. Answer: dwarf spheroidal.")
    d.insert_marker("ellipticals: [Mg/Fe]–mass, formation timescales, time-variable IMF (Yan, Jeřábková+2019, 2021)",
                    "Canva ‘ISTA seminar’, slides 52–57",
                    notes="1:52. Thomas+2005: more massive ellipticals are more α-enhanced → formed faster (downsizing). "
                          "With the IGIMF: top-heavy early (starburst) and bottom-heavy later (metal-rich) can do both.")
    d.slido("S2.9", "multiple choice", "Massive ellipticals have [Mg/Fe] ≈ +0.3. Their stars formed within …",
            ["~1 Gyr", "~5 Gyr", ">10 Gyr"], notes="2:00. Answer ~1 Gyr — with the IMF caveat.")
    d.slido("S2.10", "poll (before and after)", "Is the IMF universal?", ["yes", "no", "it depends"],
            notes="2:04–2:14. Debate: two camps, 3 min preparation, 2 min each, then vote again.")
    d.bullets("Take-aways",
              ["SN II enrich fast with α-elements; SN Ia slowly with iron.",
               "The knee in [Mg/Fe] marks when SN Ia catch up — a clock.",
               "Fast-forming galaxies (ellipticals) keep α-rich stars.",
               "The IMF sets the plateau: time and IMF are degenerate.",
               "The IMF links light → mass (Lecture 1) and chemistry → time (Lecture 2)."],
              notes="2:14")
    d.bullets("Homework & exit ticket",
              ["Notebooks 03 (toy model) and 04 (galevo)",
               "github.com/Azeret/stellar-populations-lectures   [QR code]",
               "Slido S2.11: one thing you will remember"],
              notes="2:18")
    d.save("lecture2_chemical_evolution.pptx")


if __name__ == "__main__":
    lecture1(); lecture2()
